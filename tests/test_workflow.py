"""Suppression, tracking, export, approval/sending safety, the pipeline and the CLI."""

from __future__ import annotations

import csv
import shutil
from datetime import date, timedelta
from pathlib import Path

import pytest

from src import models as m
from src.drafting import apply_draft, build_draft, draft_hash
from src.errors import UserFacingError
from src.export import _safe_cell, export_all
from src.run_pipeline import run
from src.sending import approve, mark_sent, send
from src.storage import load_records, save_records
from src.suppression import SuppressionList
from src.tracker import refresh_follow_ups, sent_count_on, set_status

from .conftest import ROOT, make_record, write_profile


def drafted(profile, today, **overrides):
    rec = make_record(**overrides)
    apply_draft(rec, build_draft(rec, profile, today), today)
    return rec


# ------------------------------------------------------------------ suppression + tracker
def test_suppression_blocks_by_email_and_by_name(settings, today):
    sup = SuppressionList()
    sup.add(make_record(), m.OPTED_OUT, today)
    sup.save(settings.suppression_path)
    sup = SuppressionList.load(settings.suppression_path)
    same_email = make_record(record_id="R0009")
    same_person_new_email = make_record(record_id="R0010", institutional_email="new@testing.edu")
    other = make_record(record_id="R0011", professor_name="Other Person", institutional_email="other@testing.edu")
    records = [same_email, same_person_new_email, other]
    assert sup.apply(records, today) == 2
    assert same_email["workflow_status"] == m.OPTED_OUT
    assert same_person_new_email["workflow_status"] == m.OPTED_OUT
    assert other["workflow_status"] == m.VERIFIED


def test_sent_schedules_one_follow_up(settings, today):
    rec = make_record(workflow_status=m.APPROVED)
    set_status(rec, m.SENT, today, settings)
    assert rec["last_contact_date"] == today.isoformat()
    assert rec["follow_up_date"] == (today + timedelta(days=settings.follow_up_days)).isoformat()
    assert refresh_follow_ups([rec], today, settings) == []
    later = today + timedelta(days=settings.follow_up_days)
    assert refresh_follow_ups([rec], later, settings) == [rec["record_id"]]
    assert rec["workflow_status"] == m.FOLLOW_UP_DUE
    set_status(rec, m.DRAFT_READY, later, settings)  # follow-up drafted
    set_status(rec, m.APPROVED, later, settings)  # and approved like any email
    set_status(rec, m.SENT, later, settings)  # the follow-up was sent
    assert rec["follow_up_count"] == "1"
    assert rec["follow_up_date"] == ""  # no second follow-up by default


# ------------------------------------------------------------------ export
def test_exports_and_formula_injection_guard(settings, profile, today):
    rec = drafted(profile, today, notes="=HYPERLINK(\"http://evil\")")
    csv_path, xlsx_path = export_all([rec], settings.exports_dir, programs=[{"university": "X", "program_name": "Y"}], stamp="test")
    assert csv_path.exists() and xlsx_path.exists()
    with open(csv_path, encoding="utf-8-sig", newline="") as handle:
        row = next(csv.DictReader(handle))
    assert row["notes"].startswith("'=")
    from openpyxl import load_workbook

    wb = load_workbook(xlsx_path)
    assert {"Summary", "All fields", "Drafts", "Programs", "Status counts"} <= set(wb.sheetnames)
    ws = wb["All fields"]
    header = [c.value for c in ws[1]]
    url_col = header.index("faculty_profile_url") + 1
    assert ws.cell(row=2, column=url_col).hyperlink is not None
    assert _safe_cell("-5") == "-5" and _safe_cell("+1") == "'+1"


# ------------------------------------------------------------------ approval + sending safety
def answers(*values):
    it = iter(values)
    return lambda prompt="": next(it)


def test_approval_requires_explicit_yes(settings, profile, today):
    rec = drafted(profile, today)
    sup = SuppressionList()
    assert approve([rec], settings, sup, today, input_fn=answers("y"), print_fn=lambda *a: None) == 0
    assert rec["workflow_status"] == m.DRAFT_READY
    assert approve([rec], settings, sup, today, input_fn=answers("yes"), print_fn=lambda *a: None) == 1
    assert rec["workflow_status"] == m.APPROVED
    assert rec["approval_hash"] == draft_hash(rec)


def test_editing_after_approval_blocks_sending(settings, profile, today):
    rec = drafted(profile, today)
    approve([rec], settings, SuppressionList(), today, input_fn=answers("yes"), print_fn=lambda *a: None)
    rec["email_draft"] += "\nP.S. edited later"
    report = send([rec], settings, SuppressionList(), today, method="manual", print_fn=lambda *a: None)
    assert not report.outbox_path
    assert "approve it again" in report.skipped[0][1]


def test_manual_send_makes_outbox_but_changes_nothing(settings, profile, today):
    rec = drafted(profile, today)
    approve([rec], settings, SuppressionList(), today, input_fn=answers("yes"), print_fn=lambda *a: None)
    report = send([rec], settings, SuppressionList(), today, method="manual", print_fn=lambda *a: None)
    assert report.outbox_path and report.outbox_path.exists()
    assert "mail.google.com" in report.outbox_path.read_text(encoding="utf-8")
    assert rec["workflow_status"] == m.APPROVED  # only mark-sent records a send


def test_gmail_api_sending_is_off_by_default(settings, profile, today):
    rec = drafted(profile, today)
    approve([rec], settings, SuppressionList(), today, input_fn=answers("yes"), print_fn=lambda *a: None)
    with pytest.raises(UserFacingError):
        send([rec], settings, SuppressionList(), today, method="gmail-api", gmail_service=object(), print_fn=lambda *a: None)


class FakeGmail:
    def __init__(self, fail_on=None):
        self.sent, self.created_drafts, self.fail_on = [], [], fail_on

    def users(self):
        return self

    def messages(self):
        return self

    def drafts(self):
        return _Drafts(self)

    def send(self, userId, body):
        return _Exec(self, "send", body)


class _Drafts:
    def __init__(self, owner):
        self.owner = owner

    def create(self, userId, body):
        return _Exec(self.owner, "draft", body)


class _Exec:
    def __init__(self, owner, kind, body):
        self.owner, self.kind, self.body = owner, kind, body

    def execute(self):
        if self.owner.fail_on is not None and len(self.owner.sent) == self.owner.fail_on:
            raise RuntimeError("network down")
        if self.kind == "send":
            self.owner.sent.append(self.body)
            return {"id": f"msg{len(self.owner.sent)}"}
        self.owner.created_drafts.append(self.body)
        return {"id": f"draft{len(self.owner.created_drafts)}"}


def _approved(settings, profile, today, n):
    recs = []
    for i in range(n):
        rec = drafted(profile, today, record_id=f"R{i + 1:04d}", institutional_email=f"prof{i}@testing.edu", professor_name=f"Prof Number{i}")
        approve([rec], settings, SuppressionList(), today, input_fn=answers("yes"), print_fn=lambda *a: None)
        recs.append(rec)
    return recs


def test_gmail_api_needs_exact_phrase_and_respects_daily_limit(settings, profile, today):
    settings.email_mode = "send_enabled"
    settings.daily_send_limit = 2
    recs = _approved(settings, profile, today, 3)
    gmail = FakeGmail()
    report = send(recs, settings, SuppressionList(), today, method="gmail-api", gmail_service=gmail, input_fn=answers("send 2"), print_fn=lambda *a: None, sleep=lambda s: None)
    assert report.sent == [] and gmail.sent == []  # wrong phrase (case matters): nothing sent
    report = send(recs, settings, SuppressionList(), today, method="gmail-api", gmail_service=gmail, input_fn=answers("SEND 2"), print_fn=lambda *a: None, sleep=lambda s: None)
    assert len(report.sent) == 2 and len(gmail.sent) == 2
    assert any("daily limit" in why for _, why in report.skipped)
    assert sent_count_on(settings, date.today()) == 2  # the outreach log is stamped with the real clock
    assert [r["workflow_status"] for r in recs].count(m.SENT) == 2


def test_failed_send_is_never_marked_sent(settings, profile, today):
    settings.email_mode = "send_enabled"
    recs = _approved(settings, profile, today, 2)
    gmail = FakeGmail(fail_on=1)
    report = send(recs, settings, SuppressionList(), today, method="gmail-api", gmail_service=gmail, input_fn=answers("SEND 2"), print_fn=lambda *a: None, sleep=lambda s: None)
    assert len(report.sent) == 1 and len(report.failed) == 1
    assert recs[1]["workflow_status"] == m.APPROVED


def test_gmail_drafts_do_not_mark_sent(settings, profile, today):
    recs = _approved(settings, profile, today, 1)
    gmail = FakeGmail()
    report = send(recs, settings, SuppressionList(), today, method="gmail-drafts", gmail_service=gmail, input_fn=answers("DRAFT 1"), print_fn=lambda *a: None)
    assert report.drafted == ["R0001"] and gmail.created_drafts and not gmail.sent
    assert recs[0]["workflow_status"] == m.APPROVED


def test_sample_and_suppressed_records_are_never_sent(settings, profile, today):
    settings.email_mode = "send_enabled"
    sample = drafted(profile, today, record_id="R0001", institutional_email="x@uni.example", record_source="sample", email_verification="found_on_official_page")
    approve([sample], settings, SuppressionList(), today, input_fn=answers("yes"), print_fn=lambda *a: None)
    report = send([sample], settings, SuppressionList(), today, method="gmail-api", gmail_service=FakeGmail(), print_fn=lambda *a: None)
    assert report.skipped and "never sent" in report.skipped[0][1]
    real = _approved(settings, profile, today, 1)[0]
    sup = SuppressionList()
    sup.add(real, m.DO_NOT_CONTACT, today)
    report = send([real], settings, sup, today, method="gmail-api", gmail_service=FakeGmail(), print_fn=lambda *a: None)
    assert "suppression" in report.skipped[0][1]


def test_mark_sent_requires_confirmation(settings, profile, today):
    rec = _approved(settings, profile, today, 1)[0]
    assert mark_sent([rec], ["R0001"], settings, today, input_fn=answers("no"), print_fn=lambda *a: None) == 0
    assert mark_sent([rec], ["R0001"], settings, today, input_fn=answers("yes"), print_fn=lambda *a: None) == 1
    assert rec["workflow_status"] == m.SENT


# ------------------------------------------------------------------ pipeline (end to end, offline)
def test_pipeline_on_sample_data(settings, profile, today):
    write_profile(settings, profile)
    shutil.copy(ROOT / "data" / "sample" / "sample_suppression_list.csv", settings.suppression_path)
    report = run(settings, intake="both", limit=30, mode="drafts", input_file=ROOT / "data" / "sample" / "sample_professors.csv", today=today, print_fn=lambda *a: None)
    records = {r["record_id"]: r for r in load_records(settings.tracker_path)}
    assert report.merged == 1  # S005 duplicates S001
    assert "S005" not in records
    assert set(report.drafted) == {"S001", "S002", "S003"}
    assert records["S004"]["workflow_status"] == m.NEEDS_VERIFICATION  # unverified: no draft
    assert records["S006"]["workflow_status"] == m.VERIFIED  # department discourages contact: no draft
    assert records["S007"]["workflow_status"] == m.OPTED_OUT
    assert not any(r["workflow_status"] == m.SENT for r in records.values())  # the pipeline never sends
    names = {p.name for p in report.files}
    assert "review_latest.html" in names and any(n.endswith(".xlsx") for n in names) and any(n.startswith("run_report_") for n in names)
    assert "S001" in (settings.exports_dir / "review_latest.html").read_text(encoding="utf-8")


def test_pipeline_refuses_send_mode(settings, profile, today):
    write_profile(settings, profile)
    with pytest.raises(UserFacingError):
        run(settings, mode="send", today=today)


def test_pipeline_rerun_is_stable(settings, profile, today):
    write_profile(settings, profile)
    sample = ROOT / "data" / "sample" / "sample_professors.csv"
    run(settings, input_file=sample, today=today, print_fn=lambda *a: None)
    first = load_records(settings.tracker_path)
    report = run(settings, today=today, print_fn=lambda *a: None)
    second = load_records(settings.tracker_path)
    assert len(first) == len(second)
    assert report.drafted == []  # existing drafts are not rewritten
    assert {r["record_id"]: r["email_draft"] for r in first} == {r["record_id"]: r["email_draft"] for r in second}


# ------------------------------------------------------------------ CLI smoke tests
def test_cli_commands(tmp_path, monkeypatch, profile, capsys):
    from src import cli
    from src.config import load_settings

    shutil.copytree(ROOT / "data" / "sample", tmp_path / "data" / "sample")
    monkeypatch.setenv("OUTREACH_ROOT", str(tmp_path))
    settings = load_settings()
    settings.ensure_folders()
    write_profile(settings, profile)
    assert cli.main(["doctor"]) == 0
    assert cli.main(["init", "--sample"]) == 0
    assert cli.main(["status"]) == 0
    assert cli.main(["score"]) == 0
    assert cli.main(["draft", "--limit", "5"]) == 0
    assert cli.main(["review"]) == 0
    assert cli.main(["validate"]) == 0
    assert cli.main(["export"]) == 0
    assert cli.main(["not-a-fit", "--id", "S002"]) == 0
    assert cli.main(["opt-out", "--id", "S003", "--note", "asked to stop"]) == 0
    sup = SuppressionList.load(settings.suppression_path)
    assert sup.is_suppressed({"institutional_email": "rplaceholder@fictional-tech.example"})
    assert cli.main(["mark-sent", "--ids", "S001"]) == 0  # not approved: politely refused, nothing changes
    out = capsys.readouterr().out
    assert "Only approved emails" in out
    assert cli.main(["confirm-email", "--id", "S004", "--email", "m@x.edu", "--source-url", "https://x.edu/p"]) == 0
    assert cli.main(["list"]) == 0
    assert cli.main(["init", "--sample"]) == 1  # refuses to overwrite without --force
    assert (Path(tmp_path) / "workspace" / "outreach_log.csv").exists()


def test_cli_reports_friendly_error_for_missing_tracker(tmp_path, monkeypatch, capsys):
    from src import cli

    monkeypatch.setenv("OUTREACH_ROOT", str(tmp_path))
    assert cli.main(["status"]) == 1
    assert "tracker is empty" in capsys.readouterr().err


def test_saving_while_file_locked_gives_clear_message(tmp_path, monkeypatch):
    import os

    rec = make_record()
    path = tmp_path / "tracker.csv"

    def locked(*args, **kwargs):
        raise PermissionError("locked by Excel")

    monkeypatch.setattr(os, "replace", locked)
    with pytest.raises(UserFacingError, match="Excel"):
        save_records(path, [rec])


def test_preview_is_read_only(tmp_path, monkeypatch, profile):
    from src import cli
    from src.config import load_settings

    shutil.copytree(ROOT / "data" / "sample", tmp_path / "data" / "sample")
    monkeypatch.setenv("OUTREACH_ROOT", str(tmp_path))
    settings = load_settings()
    settings.ensure_folders()
    write_profile(settings, profile)
    assert cli.main(["init", "--sample"]) == 0
    before = settings.tracker_path.read_bytes()
    assert cli.main(["preview", "--limit", "3"]) == 0
    assert settings.tracker_path.read_bytes() == before  # nothing saved
    page = (settings.exports_dir / "preview_latest.html").read_text(encoding="utf-8")
    assert "PREVIEW" in page and "mail.google.com" not in page  # no send buttons on a preview
