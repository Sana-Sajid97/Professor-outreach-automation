"""Export the tracker to dated CSV and Excel (XLSX) files."""

from __future__ import annotations

from datetime import datetime
from pathlib import Path

from . import models as m
from .storage import ordered_columns, write_csv
from .textutil import clean, is_valid_url

SUMMARY_COLUMNS = [
    "record_id",
    "fit_score",
    "workflow_status",
    "professor_name",
    "title",
    "university",
    "state",
    "department_or_school",
    "institutional_email",
    "faculty_profile_url",
    "research_interests",
    "research_fit_explanation",
    "graduate_program",
    "spring_2027_status",
    "fall_2027_status",
    "funding_level",
    "funding_or_assistantship_evidence",
    "verification_status",
    "checks_summary",
    "next_action",
    "follow_up_date",
    "fit_score_breakdown",
]
DRAFT_COLUMNS = ["record_id", "professor_name", "university", "institutional_email", "email_subject", "email_draft", "workflow_status",
                 "checks_summary", "draft_edited_on", "cv_fingerprint"]
WIDE = {"research_interests", "research_fit_explanation", "fit_score_breakdown", "email_draft", "notes", "funding_or_assistantship_evidence",
        "next_action", "eligibility_notes", "source_urls", "checks_summary", "verification_checks"}


def _safe_cell(value: str) -> str:
    """Stop spreadsheet programs from treating text as a formula (CSV/Excel injection)."""
    text = "" if value is None else str(value)
    if text[:1] in ("=", "+", "@", "\t", "\r") or (text[:1] == "-" and len(text) > 1 and not text[1:2].isdigit()):
        return "'" + text
    return text


def sort_for_review(records: list[dict]) -> list[dict]:
    def key(r: dict):
        try:
            score = float(clean(r.get("fit_score")) or 0)
        except ValueError:
            score = 0.0
        return (m.STATUS_RANK.get(clean(r.get("workflow_status")), 0) >= 8, -score, clean(r.get("university")), clean(r.get("professor_name")))

    return sorted(records, key=key)


def export_csv(records: list[dict], path: Path) -> Path:
    columns = ordered_columns(records)
    safe = [{c: _safe_cell(r.get(c, "")) for c in columns} for r in records]
    write_csv(path, safe, columns)
    return Path(path)


def export_xlsx(records: list[dict], path: Path, programs: list[dict] | None = None) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="2F5597")
    link_font = Font(color="0563C1", underline="single")

    def write_sheet(ws, rows: list[dict], columns: list[str]) -> None:
        ws.append(columns)
        for cell in ws[1]:
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        for row in rows:
            ws.append([_safe_cell(clean(row.get(c, ""))) for c in columns])
        for idx, column in enumerate(columns, start=1):
            letter = get_column_letter(idx)
            ws.column_dimensions[letter].width = 60 if column in WIDE else (14 if column in {"record_id", "fit_score", "state"} else 28)
            for cell in ws[letter][1:]:
                cell.data_type = "s"
                cell.alignment = Alignment(wrap_text=column in WIDE, vertical="top")
                value = str(cell.value or "")
                if column.endswith("_url") and is_valid_url(value):
                    cell.hyperlink = value
                    cell.font = link_font
        ws.freeze_panes = "B2"
        if rows:
            ws.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{len(rows) + 1}"

    ordered = sort_for_review(records)
    ws = wb.active
    ws.title = "Summary"
    write_sheet(ws, ordered, [c for c in SUMMARY_COLUMNS])
    write_sheet(wb.create_sheet("All fields"), ordered, ordered_columns(records))  # a snapshot: no dropdowns that invite edits
    drafts = [r for r in ordered if clean(r.get("email_draft"))]
    write_sheet(wb.create_sheet("Drafts"), drafts, DRAFT_COLUMNS)
    if programs:
        columns: list[str] = []
        for p in programs:
            for k in p:
                if k not in columns:
                    columns.append(k)
        write_sheet(wb.create_sheet("Programs"), programs, columns)
    counts = {s: 0 for s in m.WORKFLOW_STATUSES}
    for r in records:
        counts[clean(r.get("workflow_status")) or m.DISCOVERED] = counts.get(clean(r.get("workflow_status")) or m.DISCOVERED, 0) + 1
    ws = wb.create_sheet("Status counts")
    ws.append(["workflow_status", "count"])
    for status, count in counts.items():
        ws.append([status, count])
    ws.append([])
    ws.append(["Fit scores are a prioritisation aid only - not an admission probability."])
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    try:
        wb.save(path)
    except PermissionError:
        from .errors import UserFacingError

        raise UserFacingError(f"Could not save {path.name}: close it in Excel and try again.")
    return path


def export_all(records: list[dict], exports_dir: Path, programs: list[dict] | None = None, stamp: str | None = None) -> tuple[Path, Path]:
    stamp = stamp or datetime.now().strftime("%Y-%m-%d_%H%M")
    exports_dir = Path(exports_dir)
    csv_path = export_csv(sort_for_review(records), exports_dir / f"professors_{stamp}.csv")
    xlsx_path = export_xlsx(records, exports_dir / f"professors_{stamp}.xlsx", programs)
    return csv_path, xlsx_path
